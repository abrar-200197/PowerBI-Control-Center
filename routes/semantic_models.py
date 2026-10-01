"""Semantic Models blueprint — page + API routes.

Extracted from app.py verbatim (no behavior change).

NOTE: `login_required`, `CATALOG_AVAILABLE`, `catalog_service`, and
`get_user_powerbi_headers` live in app.py. They are intentionally NOT
imported at module level here — doing so causes a circular import when
app.py is run directly as a script (`python app.py`), because the file is
then loaded as `__main__` and `from app import ...` triggers Python to
re-import app.py fresh under the name `app`, re-running the blueprint
registration line while this module is still mid-import. Instead, resolve
them lazily at call/request time, by which point `app` is fully initialized
in sys.modules.

`_semantic_scan_cache`, `_semantic_scan_lock`, `_get_workspace_scan_cached`,
and `_extract_measures_and_relationships` are defined HERE (this is their
home module) and re-exported from app.py so that routes/similarity.py's
existing lazy proxies (`from app import _semantic_scan_lock`, etc.) keep
working unchanged.
"""
import time
import threading
from functools import wraps
from flask import Blueprint, request, jsonify, render_template, send_file

semantic_models_bp = Blueprint('semantic_models', __name__)


def login_required(f):
    """Lazy proxy for app.login_required (avoids circular import at load time)."""
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not hasattr(wrapper, '_decorated'):
            from app import login_required as _real_login_required
            wrapper._decorated = _real_login_required(f)
        return wrapper._decorated(*args, **kwargs)
    return wrapper


def _catalog_available():
    from app import CATALOG_AVAILABLE as _real
    return _real


def _get_catalog_service():
    from app import catalog_service as _real
    return _real


def _get_user_powerbi_headers():
    from app import get_user_powerbi_headers as _real
    return _real()


@semantic_models_bp.route('/semantic-models')
@login_required
def semantic_models_page():
    """Semantic Models page - Analyze and health-check semantic models"""
    return render_template('semantic_models.html')


# Workspace-level Scanner schema cache for Semantic Models Details
# (one scan fills measures/relationships for all models in that workspace)
_semantic_scan_cache = {}
_semantic_scan_lock = threading.Lock()
_SEMANTIC_SCAN_TTL_SEC = 45 * 60  # 45 minutes


def _get_workspace_scan_cached(workspace_id: str) -> dict:
    """Run Admin Scanner once per workspace and cache datasets by id."""
    now = time.time()
    with _semantic_scan_lock:
        hit = _semantic_scan_cache.get(workspace_id)
        if hit and (now - hit.get('ts', 0)) < _SEMANTIC_SCAN_TTL_SEC:
            return hit.get('by_id') or {}

    from scanner_connector import PowerBIScanner
    print(f"📡 Scanner schema scan for semantic details ws={workspace_id[:8]}…")
    scanner = PowerBIScanner()
    scan_result = scanner.run_scan(workspace_id=workspace_id) or {}
    workspaces = scan_result.get('workspaces') or []
    by_id = {}
    if workspaces:
        for ds in workspaces[0].get('datasets') or []:
            did = ds.get('id')
            if did:
                by_id[did] = ds
    print(f"   cached {len(by_id)} dataset schemas for workspace {workspace_id[:8]}…")

    with _semantic_scan_lock:
        _semantic_scan_cache[workspace_id] = {'ts': time.time(), 'by_id': by_id}
    return by_id


def _get_scanned_dataset_cached(workspace_id: str, dataset_id: str):
    by_id = _get_workspace_scan_cached(workspace_id)
    return by_id.get(dataset_id)


def _merge_dataset_schema(base: dict, scan_ds: dict) -> dict:
    """Merge scanner measures/relationships/columns onto catalog dataset."""
    out = dict(base or {})
    if not scan_ds:
        return out

    if not out.get('relationships') and scan_ds.get('relationships'):
        out['relationships'] = scan_ds.get('relationships')
    if scan_ds.get('configuredBy') and not out.get('configuredBy'):
        out['configuredBy'] = scan_ds.get('configuredBy')
    if scan_ds.get('createdDate') and not out.get('createdDate'):
        out['createdDate'] = scan_ds.get('createdDate')

    scan_tables = {
        (t.get('name') or ''): t for t in (scan_ds.get('tables') or [])
    }
    merged = []
    for t in out.get('tables') or []:
        tt = dict(t)
        st = scan_tables.get(tt.get('name') or '')
        if st:
            if not tt.get('measures') and st.get('measures'):
                tt['measures'] = st.get('measures')
                tt['measureCount'] = len(st.get('measures') or [])
            if not tt.get('columns') and st.get('columns'):
                tt['columns'] = st.get('columns')
                tt['columnCount'] = len(st.get('columns') or [])
            # fill measure expressions if catalog only had empty measures
            if tt.get('measures') and st.get('measures'):
                by_name = {
                    (m.get('name') if isinstance(m, dict) else str(m)): m
                    for m in (st.get('measures') or [])
                    if isinstance(m, dict) or isinstance(m, str)
                }
                fixed = []
                for m in tt.get('measures') or []:
                    if isinstance(m, dict) and not m.get('expression'):
                        sm = by_name.get(m.get('name') or '')
                        if isinstance(sm, dict) and sm.get('expression'):
                            mm = dict(m)
                            mm['expression'] = sm.get('expression')
                            fixed.append(mm)
                            continue
                    fixed.append(m)
                tt['measures'] = fixed
        merged.append(tt)

    have = {t.get('name') for t in merged}
    for name, st in scan_tables.items():
        if name and name not in have:
            merged.append(st)
    out['tables'] = merged
    if not out.get('relationshipCount'):
        out['relationshipCount'] = len(out.get('relationships') or [])
    return out


def _normalize_semantic_tables(tables):
    """Return UI-friendly table objects with explicit column lists."""
    out = []
    for table in tables or []:
        cols_in = table.get('columns') or []
        columns = []
        for c in cols_in:
            if isinstance(c, str):
                columns.append({'name': c, 'dataType': '', 'isHidden': False})
            elif isinstance(c, dict):
                columns.append({
                    'name': c.get('name') or c.get('columnName') or '',
                    'dataType': c.get('dataType') or c.get('dataTypeName') or c.get('type') or '',
                    'isHidden': bool(c.get('isHidden') or c.get('isHiddenColumn')),
                    'usedInReport': c.get('usedInReport'),
                })
        measures_in = table.get('measures') or []
        measures = []
        for m in measures_in:
            if isinstance(m, str):
                measures.append({'name': m, 'expression': ''})
            elif isinstance(m, dict):
                measures.append({
                    'name': m.get('name') or '',
                    'expression': m.get('expression') or '',
                })
        out.append({
            'name': table.get('name') or 'Unknown',
            'isHidden': bool(table.get('isHidden')),
            'columnCount': table.get('columnCount') if table.get('columnCount') is not None else len(columns),
            'measureCount': table.get('measureCount') if table.get('measureCount') is not None else len(measures),
            'columns': columns,
            'measures': measures,
            'sourceTypeLabel': table.get('sourceTypeLabel'),
            'serverName': table.get('serverName'),
            'sqlSourceTables': table.get('sqlSourceTables') or [],
            'sqlQuery': table.get('sqlQuery') or '',
            'fileName': table.get('fileName') or '',
            'sourceUrl': table.get('sourceUrl') or '',
            'sourceExpression': table.get('sourceExpression') or '',
        })
    return out


def _rel_field(rel, *keys):
    """Read a relationship endpoint from common Scanner / catalog key shapes."""
    if not isinstance(rel, dict):
        return ''
    for k in keys:
        v = rel.get(k)
        if v is None or v == '':
            continue
        if isinstance(v, dict):
            # nested { table, column } / { name }
            return (
                v.get('table')
                or v.get('tableName')
                or v.get('column')
                or v.get('columnName')
                or v.get('name')
                or ''
            )
        return v
    return ''


def _extract_measures_and_relationships(dataset):
    tables = dataset.get('tables') or []
    all_measures = []
    for table in tables:
        tname = table.get('name') or 'Unknown'
        for measure in table.get('measures') or []:
            if isinstance(measure, str):
                all_measures.append({'name': measure, 'table': tname, 'expression': ''})
            elif isinstance(measure, dict):
                all_measures.append({
                    'name': measure.get('name'),
                    'table': tname,
                    'expression': measure.get('expression') or '',
                })

    # Some extracts put measures only under dataset.expressions
    if not all_measures:
        for expr in dataset.get('expressions') or []:
            if not isinstance(expr, dict):
                continue
            name = expr.get('name') or ''
            body = expr.get('expression') or expr.get('query') or ''
            if name and body:
                all_measures.append({
                    'name': name,
                    'table': expr.get('table') or 'Expression',
                    'expression': body if isinstance(body, str) else str(body),
                })

    raw_rels = (
        dataset.get('relationships')
        or dataset.get('modelRelationships')
        or dataset.get('datasetRelationships')
        or []
    )
    relationships = []
    for rel in raw_rels:
        if not isinstance(rel, dict):
            continue
        from_table = _rel_field(
            rel, 'fromTable', 'sourceTable', 'fromTableName', 'from', 'FromTable',
        )
        from_col = _rel_field(
            rel, 'fromColumn', 'sourceColumn', 'fromColumnName', 'FromColumn',
        )
        to_table = _rel_field(
            rel, 'toTable', 'targetTable', 'toTableName', 'to', 'ToTable',
        )
        to_col = _rel_field(
            rel, 'toColumn', 'targetColumn', 'toColumnName', 'ToColumn',
        )
        # nested from/to objects: { from: { table, column }, to: {...} }
        if isinstance(rel.get('from'), dict) and not from_table:
            fr = rel['from']
            from_table = fr.get('table') or fr.get('tableName') or ''
            from_col = from_col or fr.get('column') or fr.get('columnName') or ''
        if isinstance(rel.get('to'), dict) and not to_table:
            to = rel['to']
            to_table = to.get('table') or to.get('tableName') or ''
            to_col = to_col or to.get('column') or to.get('columnName') or ''
        if not (from_table or to_table or from_col or to_col):
            continue
        relationships.append({
            'fromTable': from_table or '',
            'fromColumn': from_col or '',
            'toTable': to_table or '',
            'toColumn': to_col or '',
            'cardinality': (
                rel.get('cardinality')
                or rel.get('Cardinality')
                or rel.get('crossFilteringBehavior')
                or rel.get('relationshipType')
                or ''
            ),
            'isActive': rel.get('isActive', rel.get('IsActive', True)),
            'crossFilteringBehavior': rel.get('crossFilteringBehavior') or '',
        })
    return all_measures, relationships


@semantic_models_bp.route('/api/semantic-models')
@login_required
def get_semantic_models():
    """
    List semantic models for a workspace.

    Fast path: precomputed workspace_catalog (+ ops refresh fields).
    Fallback: one Scanner scan (NO live refresh fan-out — that was causing multi-minute hangs).
    """
    try:
        workspace_id = request.args.get('workspace_id')
        if not workspace_id:
            return jsonify({
                'success': False,
                'error': 'workspace_id parameter is required'
            }), 400

        # Optional: only when explicitly requested (and never the default)
        live_refresh = str(request.args.get('live_refresh', '0')).lower() in ('1', 'true', 'yes')

        def _measures_from_tables(tables):
            measures = []
            for table in tables or []:
                for measure in table.get('measures') or []:
                    expr = measure.get('expression') or ''
                    measures.append({
                        'name': measure.get('name'),
                        'table': table.get('name'),
                        'expression': expr[:100] if isinstance(expr, str) else '',
                    })
            return measures

        def _model_from_dataset(dataset, ops=None):
            tables = dataset.get('tables') or []
            relationships = dataset.get('relationships') or []
            measures = _measures_from_tables(tables)
            last_refresh = None
            refresh_status = 'Unknown'
            refresh_type = None

            ops = ops or {}
            # Catalog ops fields (preferred)
            last_refresh = (
                ops.get('last_refreshed')
                or dataset.get('last_refreshed')
                or dataset.get('lastRefresh')
                or dataset.get('lastRefreshTime')
            )
            refresh_status = (
                ops.get('last_refresh_status')
                or dataset.get('last_refresh_status')
                or dataset.get('refreshStatus')
                or ('Unknown' if not last_refresh else 'Completed')
            )
            refresh_type = ops.get('refresh_type') or dataset.get('refresh_type')

            return {
                'id': dataset.get('id'),
                'name': dataset.get('name'),
                'tables': tables,
                'relationships': relationships,
                'measures': measures,
                'tableCount': len(tables),
                'relationshipCount': len(relationships),
                'measureCount': len(measures),
                'lastRefresh': last_refresh,
                'refreshStatus': refresh_status,
                'refreshType': refresh_type,
                'configuredBy': dataset.get('configuredBy') or dataset.get('configuredByUser') or 'Unknown',
                'source': 'catalog' if ops is not None or dataset.get('_from_catalog') else 'scanner',
            }

        # ---------- Catalog fast path ----------
        catalog_service = _get_catalog_service()
        if _catalog_available() and catalog_service is not None and catalog_service.is_available():
            try:
                cat = catalog_service.get_workspace_catalog()
                ws = None
                if cat:
                    ws = next(
                        (w for w in (cat.get('workspaces') or []) if w.get('id') == workspace_id),
                        None,
                    )
                if ws is not None:
                    datasets = list(ws.get('datasets') or [])
                    # Some catalogs only embed dataset ids on reports — pull from top-level map
                    if not datasets and isinstance(cat.get('datasets'), dict):
                        for did, d in cat['datasets'].items():
                            if (d or {}).get('workspaceId') == workspace_id:
                                dd = dict(d)
                                dd.setdefault('id', did)
                                datasets.append(dd)

                    # Preload ops refresh snapshot once (dataset id → refresh fields)
                    snap_map = {}
                    try:
                        snap = catalog_service.get_json('refresh_snapshot.json') or {}
                        snap_map = snap.get('datasets') or {}
                    except Exception:
                        snap_map = {}

                    # Infer light stats from reports when dataset schema not embedded
                    report_ds_stats = {}
                    for r in (ws.get('reports') or []):
                        did = r.get('datasetId')
                        if not did:
                            continue
                        st = report_ds_stats.setdefault(did, {
                            'reportCount': 0,
                            'last_refreshed': r.get('last_refreshed'),
                            'last_refresh_status': r.get('last_refresh_status'),
                            'refresh_type': r.get('refresh_type'),
                        })
                        st['reportCount'] += 1
                        # prefer report ops if present
                        if r.get('last_refreshed') and not st.get('last_refreshed'):
                            st['last_refreshed'] = r.get('last_refreshed')
                        if r.get('last_refresh_status') and not st.get('last_refresh_status'):
                            st['last_refresh_status'] = r.get('last_refresh_status')

                    # Full schema lives on top-level cat['datasets'] map (not always on ws.datasets)
                    dmap = cat.get('datasets') if isinstance(cat.get('datasets'), dict) else {}

                    models = []
                    for ds in datasets:
                        ds = dict(ds or {})
                        ds['_from_catalog'] = True
                        did = ds.get('id') or ''
                        rich = dict(dmap.get(did) or {}) if did else {}
                        info = snap_map.get(did) or {}
                        rstat = report_ds_stats.get(did) or {}
                        ops = {
                            'last_refreshed': (
                                ds.get('last_refreshed')
                                or rich.get('last_refreshed')
                                or info.get('last_refreshed')
                                or rstat.get('last_refreshed')
                            ),
                            'last_refresh_status': (
                                ds.get('last_refresh_status')
                                or rich.get('last_refresh_status')
                                or info.get('last_refresh_status')
                                or rstat.get('last_refresh_status')
                            ),
                            'refresh_type': (
                                ds.get('refresh_type')
                                or rich.get('refresh_type')
                                or info.get('refresh_type')
                                or rstat.get('refresh_type')
                            ),
                            'refresh_schedule': (
                                ds.get('refresh_schedule')
                                or rich.get('refresh_schedule')
                                or info.get('refresh_schedule')
                            ),
                        }
                        model = _model_from_dataset(ds, ops=ops)

                        # Counts from rich schema when list entry is thin
                        rich_tables = rich.get('tables') or []
                        table_count = (
                            ds.get('tableCount')
                            or rich.get('tableCount')
                            or len(rich_tables)
                            or 0
                        )
                        measure_count = 0
                        for t in rich_tables:
                            measure_count += len(t.get('measures') or [])
                            if t.get('measureCount'):
                                # prefer explicit count if measures array empty in extract
                                if not t.get('measures'):
                                    measure_count += int(t.get('measureCount') or 0)
                        rel_count = len(rich.get('relationships') or ds.get('relationships') or [])

                        model['tableCount'] = table_count
                        model['measureCount'] = measure_count
                        model['relationshipCount'] = rel_count
                        model['reportCount'] = rstat.get('reportCount') or 0
                        # List payload stays light — full schema on Details
                        model['tables'] = []
                        model['relationships'] = []
                        model['measures'] = []
                        models.append(model)

                    models.sort(key=lambda m: (m.get('name') or '').lower())
                    print(
                        f"⚡ SEMANTIC MODELS from catalog: ws={workspace_id[:8]}… "
                        f"models={len(models)} opsHit={sum(1 for m in models if m.get('lastRefresh'))} "
                        f"opsEnrichedAt={cat.get('opsEnrichedAt')}"
                    )
                    return jsonify({
                        'success': True,
                        'models': models,
                        'source': 'catalog',
                        'opsEnrichedAt': cat.get('opsEnrichedAt'),
                        'generatedAt': cat.get('generatedAt'),
                    })
            except Exception as cat_err:
                print(f"⚠️ Semantic models catalog path failed, falling back to scanner: {cat_err}")

        # ---------- Scanner fallback (list only; no per-dataset live refresh) ----------
        from scanner_connector import PowerBIScanner

        print(f"📊 Fetching semantic models via Scanner for workspace {workspace_id}...")
        scanner = PowerBIScanner()
        scan_result = scanner.run_scan(workspace_id=workspace_id)
        if not scan_result:
            return jsonify({'success': False, 'error': 'Failed to scan workspace'}), 500

        workspaces = scan_result.get('workspaces', [])
        if not workspaces:
            return jsonify({'success': True, 'models': [], 'source': 'scanner'})

        datasets = workspaces[0].get('datasets', []) or []
        print(f"✅ Scanner returned {len(datasets)} semantic model(s)")

        # Optional live refresh (opt-in only) — never default; kills the UI spinner for minutes
        refresh_by_id = {}
        if live_refresh and datasets:
            print(f"⚠️ live_refresh=1 — resolving refresh for {len(datasets)} models (slow)")
            try:
                from powerbi_connector import resolve_dataset_refresh_info
                headers = _get_user_powerbi_headers()
                for dataset in datasets:
                    did = dataset.get('id')
                    if not did:
                        continue
                    try:
                        refresh_by_id[did] = resolve_dataset_refresh_info(
                            headers=headers,
                            workspace_id=workspace_id,
                            dataset_id=did,
                            dataset_workspace_id=workspace_id,
                            history_top=3,
                            timeout=6,
                        )
                    except Exception:
                        continue
            except Exception as e:
                print(f"⚠️ live refresh batch failed: {e}")

        models = []
        for dataset in datasets:
            ops = None
            info = refresh_by_id.get(dataset.get('id') or '')
            if info:
                ops = {
                    'last_refreshed': info.get('last_refreshed'),
                    'last_refresh_status': info.get('last_refresh_status'),
                    'refresh_type': info.get('refresh_type'),
                }
            models.append(_model_from_dataset(dataset, ops=ops))

        models.sort(key=lambda m: (m.get('name') or '').lower())
        print(f"✅ Returning {len(models)} semantic model(s) (source=scanner, live_refresh={live_refresh})")
        return jsonify({
            'success': True,
            'models': models,
            'source': 'scanner',
            'liveRefresh': live_refresh,
        })

    except Exception as e:
        print(f"❌ Error fetching semantic models: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@semantic_models_bp.route('/api/semantic-model-details')
@login_required
def get_semantic_model_details():
    """
    Detailed semantic model schema for the Details modal.
    Prefer catalog datasets map (has columns) — Scanner only if missing.
    """
    try:
        workspace_id = request.args.get('workspace_id')
        dataset_id = request.args.get('dataset_id')

        if not workspace_id or not dataset_id:
            return jsonify({
                'success': False,
                'error': 'workspace_id and dataset_id are required'
            }), 400

        dataset = None
        source = None

        catalog_service = _get_catalog_service()
        # 1) Catalog top-level datasets map (full schema when extract included it)
        if _catalog_available() and catalog_service is not None and catalog_service.is_available():
            try:
                cat = catalog_service.get_workspace_catalog()
                dmap = (cat or {}).get('datasets') or {}
                if isinstance(dmap, dict) and dataset_id in dmap:
                    dataset = dict(dmap.get(dataset_id) or {})
                    dataset.setdefault('id', dataset_id)
                    source = 'catalog'
                # Also try workspace.datasets if richer
                if dataset is not None and not (dataset.get('tables') or []):
                    ws = next(
                        (w for w in (cat.get('workspaces') or []) if w.get('id') == workspace_id),
                        None,
                    )
                    if ws:
                        for ds in ws.get('datasets') or []:
                            if ds.get('id') == dataset_id and (ds.get('tables') or []):
                                dataset = dict(ds)
                                source = 'catalog-ws'
                                break
            except Exception as cat_err:
                print(f"⚠️ semantic-model-details catalog miss: {cat_err}")

        # 2) Fill missing measures / relationships / columns via Scanner (cached per workspace).
        # IMPORTANT: having measures does NOT mean relationships are present.
        # Old logic treated (has measures) as complete and skipped Scanner → Relationships (0).
        def _schema_incomplete(ds_obj):
            if not ds_obj:
                return True
            tables = ds_obj.get('tables') or []
            if not tables:
                return True
            has_cols = any((t.get('columns') or t.get('columnCount')) for t in tables)
            has_meas = (
                any((t.get('measures') or t.get('measureCount')) for t in tables)
                or bool(ds_obj.get('measureCount'))
                or bool(ds_obj.get('expressions'))
            )
            rels = (
                ds_obj.get('relationships')
                or ds_obj.get('modelRelationships')
                or []
            )
            rel_n = len(rels) if isinstance(rels, list) else 0
            if not rel_n:
                try:
                    rel_n = int(ds_obj.get('relationshipCount') or 0)
                except Exception:
                    rel_n = 0
            has_rel = rel_n > 0
            # Multi-table models almost always have relationships in Desktop —
            # if catalog has 2+ tables and 0 rels, force Scanner enrich.
            multi_table = len(tables) >= 2
            missing_rels = (not has_rel) and multi_table
            return (not has_cols) or (not has_meas) or missing_rels

        # opt-out: enrich=0 skips Scanner (catalog-only)
        allow_enrich = str(request.args.get('enrich', '1')).lower() not in ('0', 'false', 'no')
        need_scan = _schema_incomplete(dataset)

        if allow_enrich and need_scan:
            try:
                print(
                    f"🔎 semantic-model-details enriching via Scanner "
                    f"(ws={workspace_id[:8]}… ds={dataset_id[:8]}… "
                    f"had_tables={len((dataset or {}).get('tables') or [])} "
                    f"had_rels={len((dataset or {}).get('relationships') or [])})"
                )
                scan_ds = _get_scanned_dataset_cached(workspace_id, dataset_id)
                if scan_ds:
                    scan_rel_n = len(scan_ds.get('relationships') or [])
                    print(f"   scanner dataset keys sample rels={scan_rel_n}")
                    if dataset is None or not (dataset.get('tables') or []):
                        dataset = scan_ds
                        source = 'scanner'
                    else:
                        dataset = _merge_dataset_schema(dataset, scan_ds)
                        source = 'catalog+scanner'
            except Exception as scan_err:
                print(f"⚠️ semantic-model-details scanner enrich failed: {scan_err}")
                if dataset is None:
                    return jsonify({'success': False, 'error': f'Failed to load model schema: {scan_err}'}), 500

        if not dataset:
            return jsonify({
                'success': False,
                'error': f'Dataset {dataset_id} not found in workspace'
            }), 404

        raw_tables = dataset.get('tables') or []
        tables = _normalize_semantic_tables(raw_tables)
        all_measures, relationships = _extract_measures_and_relationships(dataset)

        # If still no relationships after merge, try one more direct read of scan keys
        if not relationships and allow_enrich:
            try:
                scan_ds = _get_scanned_dataset_cached(workspace_id, dataset_id)
                if scan_ds and (scan_ds.get('relationships') or []):
                    _, relationships = _extract_measures_and_relationships(scan_ds)
                    if relationships and source and 'scanner' not in str(source):
                        source = f'{source}+rels'
            except Exception:
                pass

        owner = dataset.get('configuredBy') or dataset.get('configuredByUser') or 'Unknown'
        created_date = dataset.get('createdDate') or dataset.get('createdDateTime') or 'Unknown'

        # Attach ops refresh if present
        last_refresh = dataset.get('last_refreshed') or dataset.get('lastRefresh')
        refresh_status = dataset.get('last_refresh_status') or dataset.get('refreshStatus')
        if _catalog_available() and catalog_service is not None and (not last_refresh or not refresh_status):
            try:
                snap = catalog_service.get_json('refresh_snapshot.json') or {}
                info = (snap.get('datasets') or {}).get(dataset_id) or {}
                last_refresh = last_refresh or info.get('last_refreshed')
                refresh_status = refresh_status or info.get('last_refresh_status')
            except Exception:
                pass

        print(
            f"✅ semantic-model-details source={source} tables={len(tables)} "
            f"cols={sum(len(t.get('columns') or []) for t in tables)} "
            f"measures={len(all_measures)} rels={len(relationships)}"
        )

        return jsonify({
            'success': True,
            'source': source or 'unknown',
            'tables': tables,
            'relationships': relationships,
            'measures': all_measures,
            'owner': owner,
            'configuredBy': owner,
            'modifiedBy': dataset.get('modifiedBy') or 'Not available for datasets',
            'createdDate': created_date,
            'modifiedDate': dataset.get('modifiedDateTime') or dataset.get('modifiedDate') or 'Not available for datasets',
            'name': dataset.get('name', 'Unknown'),
            'lastRefresh': last_refresh,
            'refreshStatus': refresh_status,
            'tableCount': len(tables),
            'columnCount': sum(len(t.get('columns') or []) for t in tables),
            'measureCount': len(all_measures),
            'relationshipCount': len(relationships),
        })

    except Exception as e:
        print(f"❌ Error getting semantic model details: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@semantic_models_bp.route('/api/semantic-model-health-check', methods=['POST'])
@login_required
def semantic_model_health_check():
    """Run health check on a semantic model to detect issues"""
    try:
        data = request.get_json()
        workspace_id = data.get('workspace_id')
        dataset_id = data.get('dataset_id')

        if not workspace_id or not dataset_id:
            return jsonify({
                'success': False,
                'error': 'workspace_id and dataset_id are required'
            }), 400

        from scanner_connector import PowerBIScanner

        # Get scanner instance
        scanner = PowerBIScanner()
        scan_result = scanner.run_scan(workspace_id=workspace_id)

        if not scan_result:
            return jsonify({
                'success': False,
                'error': 'Failed to scan workspace'
            }), 500

        # Extract datasets
        workspaces = scan_result.get('workspaces', [])
        if not workspaces:
            return jsonify({'success': False, 'error': 'No workspace data'}), 404

        datasets = workspaces[0].get('datasets', [])
        dataset = None
        for ds in datasets:
            if ds.get('id') == dataset_id:
                dataset = ds
                break

        if not dataset:
            return jsonify({'success': False, 'error': 'Dataset not found'}), 404

        # Perform health checks
        issues = []
        warnings = []
        health_score = 100

        tables = dataset.get('tables', [])
        relationships = dataset.get('relationships', [])

        # Check 1: Empty tables (no columns)
        for table in tables:
            if not table.get('columns') or len(table.get('columns', [])) == 0:
                issues.append({
                    'severity': 'warning',
                    'message': f"Table '{table.get('name')}' has no columns"
                })
                health_score -= 5

        # Check 2: Tables with no relationships
        table_names = [t.get('name') for t in tables]
        tables_in_relationships = set()
        for rel in relationships:
            tables_in_relationships.add(rel.get('fromTable'))
            tables_in_relationships.add(rel.get('toTable'))

        orphaned_tables = [t for t in table_names if t not in tables_in_relationships]
        if len(orphaned_tables) > 0 and len(tables) > 1:
            warnings.append(f"{len(orphaned_tables)} table(s) have no relationships: {', '.join(orphaned_tables[:3])}")
            health_score -= 10

        # Check 3: Circular relationships (simplified check)
        if len(relationships) > len(tables):
            warnings.append("Model has many relationships - check for potential circular dependencies")

        # Check 4: Tables with many columns (performance concern)
        for table in tables:
            column_count = len(table.get('columns', []))
            if column_count > 100:
                warnings.append(f"Table '{table.get('name')}' has {column_count} columns (performance concern)")
                health_score -= 5

        # Check 5: Measures without expressions
        for table in tables:
            for measure in table.get('measures', []):
                if not measure.get('expression'):
                    issues.append({
                        'severity': 'danger',
                        'message': f"Measure '{measure.get('name')}' in table '{table.get('name')}' has no expression"
                    })
                    health_score -= 10

        # Ensure score doesn't go below 0
        health_score = max(0, health_score)

        return jsonify({
            'success': True,
            'health_score': health_score,
            'issues': issues,
            'warnings': warnings,
            'summary': {
                'total_tables': len(tables),
                'total_relationships': len(relationships),
                'total_measures': sum(len(t.get('measures', [])) for t in tables),
                'orphaned_tables': len(orphaned_tables)
            }
        })

    except Exception as e:
        print(f"❌ Error running health check: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@semantic_models_bp.route('/api/semantic-model-documentation')
@login_required
def semantic_model_documentation():
    """Generate and download documentation for a semantic model"""
    try:
        workspace_id = request.args.get('workspace_id')
        dataset_id = request.args.get('dataset_id')

        if not workspace_id or not dataset_id:
            return jsonify({
                'success': False,
                'error': 'workspace_id and dataset_id are required'
            }), 400

        from scanner_connector import PowerBIScanner
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        import io

        # Get scanner instance
        scanner = PowerBIScanner()
        scan_result = scanner.run_scan(workspace_id=workspace_id)

        if not scan_result:
            return jsonify({'success': False, 'error': 'Failed to scan'}), 500

        # Extract dataset
        workspaces = scan_result.get('workspaces', [])
        if not workspaces:
            return jsonify({'success': False, 'error': 'No workspace data'}), 404

        datasets = workspaces[0].get('datasets', [])
        dataset = None
        for ds in datasets:
            if ds.get('id') == dataset_id:
                dataset = ds
                break

        if not dataset:
            return jsonify({'success': False, 'error': 'Dataset not found'}), 404

        # Create Word document
        doc = Document()

        # Title
        title = doc.add_heading('Semantic Model Documentation', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Model name
        dataset_name = dataset.get('name', 'Unknown Model')
        doc.add_heading(f"Model: {dataset_name}", 1)

        # Summary
        doc.add_heading('Summary', 2)
        tables = dataset.get('tables', [])
        relationships = dataset.get('relationships', [])

        doc.add_paragraph(f"Total Tables: {len(tables)}")
        doc.add_paragraph(f"Total Relationships: {len(relationships)}")

        total_measures = sum(len(t.get('measures', [])) for t in tables)
        doc.add_paragraph(f"Total Measures: {total_measures}")

        # Tables section
        doc.add_heading('Tables', 2)

        for table in tables:
            doc.add_heading(table.get('name', 'Unknown'), 3)

            # Columns
            columns = table.get('columns', [])
            if columns:
                doc.add_paragraph('Columns:', style='Heading 4')
                for col in columns:
                    col_text = f"  • {col.get('name', 'Unknown')} ({col.get('dataType', 'Unknown')})"
                    doc.add_paragraph(col_text)

            # Measures
            measures = table.get('measures', [])
            if measures:
                doc.add_paragraph('Measures:', style='Heading 4')
                for measure in measures:
                    doc.add_paragraph(f"  • {measure.get('name', 'Unknown')}")
                    if measure.get('expression'):
                        doc.add_paragraph(f"    Expression: {measure['expression'][:200]}")

        # Relationships section
        if relationships:
            doc.add_heading('Relationships', 2)

            for rel in relationships:
                rel_text = f"{rel.get('fromTable')}.{rel.get('fromColumn')} → {rel.get('toTable')}.{rel.get('toColumn')}"
                doc.add_paragraph(f"  • {rel_text}")
                if rel.get('crossFilteringBehavior'):
                    doc.add_paragraph(f"    Filtering: {rel['crossFilteringBehavior']}")

        # Save to BytesIO
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        return send_file(
            file_stream,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name=f'{dataset_name}_SemanticModel_Documentation.docx'
        )

    except Exception as e:
        print(f"❌ Error generating documentation: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500
